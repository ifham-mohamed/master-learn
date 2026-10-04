"""Import the supplied workbook without modifying it. Requires openpyxl."""
from pathlib import Path
from datetime import date, datetime
from collections import Counter
from html.parser import HTMLParser
import hashlib
import json
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'sources'
WORKBOOK = SOURCE / 'Final Tracker.xlsx'
if not WORKBOOK.exists():
    WORKBOOK = ROOT / 'Final Tracker.xlsx'
HTML_DIR = SOURCE / 'html' if (SOURCE / 'html').exists() else ROOT / 'Final Tracker'

def clean(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return value if value is not None else ''

class SheetParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self.row, self.cell = [], [], None
    def handle_starttag(self, tag, attrs):
        if tag == 'tr': self.row = []
        if tag == 'td':
            self.cell = None if 'freezebar-cell' in dict(attrs).get('class', '') else ''
        if tag == 'br' and self.cell is not None: self.cell += '\n'
    def handle_data(self, data):
        if self.cell is not None: self.cell += data
    def handle_endtag(self, tag):
        if tag == 'td' and self.cell is not None:
            self.row.append(self.cell)
            self.cell = None
        if tag == 'tr' and any(self.row): self.rows.append(self.row)

def records(sheet, header, keys):
    return [dict(zip(keys, map(clean, row))) | {'sourceRow': index}
            for index, row in enumerate(sheet.iter_rows(min_row=header+1, values_only=True), header+1)
            if row[0] is not None]

wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
formula_wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
task_keys = 'id startWeek category workType title priority status evidence nextAction dueDate scope estimate actual confidence doneWhen prompt prerequisites completedOn reviewDue aiMode origin originalRows sourceIds endWeek section technical communication mockResult originalSchedule completionCheck categoryKey'.split()
practice_keys = 'id introWeek practiceType category title answer priority status attempts target targetProgress confidence lastPracticed reviewDue evidence nextAction aiMode timebox easy medium hard originalRows sourceIds reviewPlan inputCheck'.split()
week_keys = 'week phase starts capacity coreEstimate optionalEstimate actual coreTasks complete completion spare dsaTarget mocks attempted outcome adjustment loadCheck focus originalDsaTarget'.split()
guide_keys = 'type title guidance relatedIds source verification'.split()
tasks = records(wb['Master Plan'], 5, task_keys)
practice = records(wb['Practice Bank'], 7, practice_keys)
weeks = [r for r in records(wb['Weekly Review'], 9, week_keys) if isinstance(r['week'], (int, float))]
guide = records(wb['Guide & Sources'], 7, guide_keys)
data = {'settings': {'startDate': clean(wb['Weekly Review']['C4'].value), 'reviewInterval': wb['Weekly Review']['C5'].value}, 'tasks': tasks, 'practice': practice, 'weeks': weeks, 'guide': guide}
out = ROOT / 'src/data'
out.mkdir(parents=True, exist_ok=True)
(out / 'tracker.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')

archive = {}
for sheet in wb:
    archive[sheet.title] = {'rows': [[clean(c) for c in row] for row in sheet.values], 'formulas': {cell.coordinate: (cell.value if isinstance(cell.value, str) else {'text': cell.value.text, 'ref': cell.value.ref}) for row in formula_wb[sheet.title] for cell in row if cell.data_type == 'f'}}
(ROOT / 'docs/workbook-archive.json').write_text(json.dumps(archive, ensure_ascii=False, indent=2), encoding='utf-8')

inventory = []
for path in [WORKBOOK, *sorted(HTML_DIR.rglob('*'))]:
    if not path.is_file(): continue
    info = {'path': str(path.relative_to(ROOT)).replace('\\', '/'), 'bytes': path.stat().st_size, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
    if path.suffix == '.html':
        parser = SheetParser()
        parser.feed(path.read_text(encoding='utf-8'))
        info['nonemptyHtmlRows'] = len(parser.rows)
        info['workbookSheet'] = path.stem if path.stem in wb.sheetnames else None
        info['htmlTaskIds'] = [row[0] for row in parser.rows if row[0] in {t['id'] for t in tasks}]
        if path.stem in Counter(t['category'] for t in tasks):
            info['matchesMasterTaskIds'] = info['htmlTaskIds'] == [t['id'] for t in tasks if t['category'] == path.stem]
    inventory.append(info)
ids = {t['id'] for t in tasks}
source_ids = {r['relatedIds'] for r in guide if r['type'] == 'Source'}
issues = []
for t in tasks:
    for ref in t['sourceIds'].split():
        if ref not in source_ids: issues.append(f"{t['id']}: unresolved source {ref}")
    if '\ufffd' in json.dumps(t, ensure_ascii=False): issues.append(f"{t['id']}: inherited replacement character (preserved)")
report = {'inventory': inventory, 'counts': {'tasks': len(tasks), 'practice': len(practice), 'weeks': len(weeks), 'sources': len(source_ids), 'categories': dict(Counter(t['category'] for t in tasks)), 'statuses': dict(Counter(t['status'] for t in tasks))}, 'issues': issues}
(ROOT / 'docs/source-audit.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
print(json.dumps(report['counts'], indent=2))
print(f'Preserved {len(archive)} sheets; audited {len(inventory)} files; {len(issues)} source observations.')
